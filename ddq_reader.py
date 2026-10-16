"""Extracts questions from a DDQ attachment (CLAUDE.md section 3/9).

Excel (.xlsx): one table, header row (possibly after title rows), one
question per row -- matches demo_inbox/DDQ_Northgate_Pension_2026.xlsx.

Word (.docx), two layouts, both real investor formats, not guessed:
(a) table layout -- same Ref/Section/Question/Response columns as Excel,
    in a Word table (matches DDQ_Harbourline_Family_Office_2026_WORD.docx);
(b) numbered-paragraph layout -- no table; section names are Word Heading
    paragraphs, and each question is a paragraph starting "1.", "Q1.",
    "Q1" etc. (see test/test_ddq_reader.py for a built sample).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import docx
import openpyxl

ID_HEADERS = {"ref", "id", "question ref", "#"}
SECTION_HEADERS = {"section"}
QUESTION_HEADERS = {"question"}
RESPONSE_HEADERS = {"manager response", "response", "answer"}

# Investor-style DDQs put a title line above the questions, e.g. "Northgate
# Pension Scheme — Manager Due Diligence Questionnaire (fictional)" -- the
# investor name is the text before the em dash (confirmed against the real
# Northgate and Harbourline files, both Excel and Word).
_TITLE_EM_DASH = "—"

_NUMBERED_QUESTION = re.compile(r"^\s*(?:Q\s*)?(\d+)[\.\)]?\s+(.+)$", re.IGNORECASE)


@dataclass
class Question:
    id: str
    section: Optional[str]
    text: str
    row: int  # 1-based worksheet row, or a paragraph/table-row index for Word


@dataclass
class DDQ:
    path: Path
    sheet_name: str  # a short label for where the questions came from (a sheet, or the Word layout)
    header_row: int
    questions: list[Question]
    investor_name: Optional[str] = None  # from the title line, if present (see investor_name_from())


def _investor_name_from_title_lines(candidate_lines: list[str]) -> Optional[str]:
    """The text before the first em dash in any candidate line, e.g.
    'Harbourline Family Office — Short-Form Manager Due Diligence
    Questionnaire (fictional)' -> 'Harbourline Family Office'."""
    for line in candidate_lines:
        if line and _TITLE_EM_DASH in line:
            name = line.split(_TITLE_EM_DASH)[0].strip()
            if name:
                return name
    return None


def _investor_name_from_filename(path: Path) -> Optional[str]:
    """E.g. 'DDQ_Harbourline_Family_Office_2026.xlsx' -> 'Harbourline
    Family Office'. The fallback when no title line matches the em-dash
    pattern."""
    stem = re.sub(r"^DDQ[_ ]+", "", path.stem, flags=re.IGNORECASE)
    stem = re.sub(r"[_ ]+(19|20)\d{2}$", "", stem)  # a trailing year, e.g. "_2026"
    name = stem.replace("_", " ").strip()
    return name or None


def investor_name_from_title(candidate_lines: list[str], path: Path) -> Optional[str]:
    """The investor name for the email template and the Word response's
    cover page (never guessed by the model -- the user's own instruction
    that the reply is 'generated from a template, not by the model')."""
    return _investor_name_from_title_lines(candidate_lines) or _investor_name_from_filename(path)


# ── Excel ────────────────────────────────────────────────────────────────


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    """Scans the first rows for the header row: investor DDQs commonly
    place one or more title rows above it (section 9's 'title rows, then a
    header row'), so this looks for the first row that has both an id-like
    and a question-like column, rather than assuming row 1."""
    for row_idx in range(1, min(ws.max_row, 10) + 1):
        col_map: dict[str, int] = {}
        for cell in ws[row_idx]:
            if cell.value is None:
                continue
            key = str(cell.value).strip().lower()
            if key in ID_HEADERS and "id" not in col_map:
                col_map["id"] = cell.column
            elif key in QUESTION_HEADERS and "question" not in col_map:
                col_map["question"] = cell.column
            elif key in SECTION_HEADERS and "section" not in col_map:
                col_map["section"] = cell.column
            elif key in RESPONSE_HEADERS and "response" not in col_map:
                col_map["response"] = cell.column
        if "id" in col_map and "question" in col_map:
            return row_idx, col_map
    raise ValueError(
        f"Could not find a header row with an id/'Ref' column and a 'Question' "
        f"column in the first 10 rows of sheet {ws.title!r}"
    )


def read_excel_ddq(path: str | Path) -> DDQ:
    """Reads one DDQ from an .xlsx attachment: finds the header row (which
    may not be row 1), then every subsequent row with both an id and a
    question becomes one Question. A 'Manager response' (or similarly
    named) column must be present -- confirms this is a DDQ table, not an
    unrelated spreadsheet (its position is no longer needed: the reply is
    always a fresh Word document now, never filled back into this file)."""
    path = Path(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook.active

    header_row, col_map = _find_header_row(worksheet)
    if "response" not in col_map:
        raise ValueError(
            f"Could not find a response column (e.g. 'Manager response') in {path.name}"
        )

    questions: list[Question] = []
    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        id_value = worksheet.cell(row=row_idx, column=col_map["id"]).value
        question_value = worksheet.cell(row=row_idx, column=col_map["question"]).value
        if id_value is None or question_value is None:
            continue  # a blank spacer row
        section_value = (
            worksheet.cell(row=row_idx, column=col_map["section"]).value
            if "section" in col_map
            else None
        )
        questions.append(
            Question(
                id=str(id_value).strip(),
                section=str(section_value).strip() if section_value is not None else None,
                text=str(question_value).strip(),
                row=row_idx,
            )
        )

    title_lines = [
        str(cell.value) for row in worksheet.iter_rows(min_row=1, max_row=header_row - 1)
        for cell in row if cell.value
    ]
    return DDQ(
        path=path,
        sheet_name=worksheet.title,
        header_row=header_row,
        questions=questions,
        investor_name=investor_name_from_title(title_lines, path),
    )


# ── Word: table layout ──────────────────────────────────────────────────


def _read_docx_table_layout(document, path: Path) -> Optional[DDQ]:
    """Layout (a): the same Ref/Section/Question/Response columns as
    Excel, in a Word table (DDQ_Harbourline_Family_Office_2026_WORD.docx)."""
    for table in document.tables:
        if not table.rows:
            continue
        header_cells = [c.text.strip().lower() for c in table.rows[0].cells]
        col_map: dict[str, int] = {}
        for idx, key in enumerate(header_cells):
            if key in ID_HEADERS and "id" not in col_map:
                col_map["id"] = idx
            elif key in QUESTION_HEADERS and "question" not in col_map:
                col_map["question"] = idx
            elif key in SECTION_HEADERS and "section" not in col_map:
                col_map["section"] = idx
            elif key in RESPONSE_HEADERS and "response" not in col_map:
                col_map["response"] = idx
        if not ("id" in col_map and "question" in col_map and "response" in col_map):
            continue  # not the DDQ table -- a document can contain other tables

        questions: list[Question] = []
        for row_idx, row in enumerate(table.rows[1:], start=1):
            cells = [c.text.strip() for c in row.cells]
            id_value = cells[col_map["id"]] if col_map["id"] < len(cells) else ""
            question_value = cells[col_map["question"]] if col_map["question"] < len(cells) else ""
            if not id_value or not question_value:
                continue
            section_value = (
                cells[col_map["section"]]
                if "section" in col_map and col_map["section"] < len(cells)
                else None
            )
            questions.append(
                Question(
                    id=id_value, section=section_value or None, text=question_value, row=row_idx,
                )
            )
        if questions:
            title_lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            return DDQ(
                path=path, sheet_name="Word (table)", header_row=0, questions=questions,
                investor_name=investor_name_from_title(title_lines, path),
            )
    return None


# ── Word: numbered-paragraph layout ─────────────────────────────────────


def _read_docx_numbered_layout(document, path: Path) -> Optional[DDQ]:
    """Layout (b): no table -- section names are Word Heading paragraphs,
    and each question is a paragraph starting '1.', 'Q1.', 'Q1 ' etc."""
    questions: list[Question] = []
    current_section: Optional[str] = None
    title_lines: list[str] = []

    for para_idx, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        is_heading = paragraph.style.name.lower().startswith("heading")
        if is_heading:
            current_section = text
            title_lines.append(text)
            continue
        match = _NUMBERED_QUESTION.match(text)
        if match:
            number, question_text = match.groups()
            questions.append(
                Question(
                    id=f"Q{int(number):02d}", section=current_section,
                    text=question_text.strip(), row=para_idx,
                )
            )
        elif not questions:
            title_lines.append(text)  # only before the first question: candidate title text

    if not questions:
        return None
    return DDQ(
        path=path, sheet_name="Word (numbered paragraphs)", header_row=0, questions=questions,
        investor_name=investor_name_from_title(title_lines, path),
    )


def read_word_ddq(path: str | Path) -> DDQ:
    """Tries the table layout first, then the numbered-paragraph layout --
    both are real investor formats, so neither is "preferred", but a
    document can only sensibly match one."""
    path = Path(path)
    document = docx.Document(str(path))

    ddq = _read_docx_table_layout(document, path)
    if ddq is not None:
        return ddq
    ddq = _read_docx_numbered_layout(document, path)
    if ddq is not None:
        return ddq
    raise ValueError(
        f"Could not find a DDQ in {path.name}: no table with Ref/Question/Response columns, "
        f"and no numbered-paragraph questions (e.g. '1.', 'Q1.') under a heading."
    )


def read_ddq(path: str | Path) -> DDQ:
    """Dispatches by file extension (section 9: 'reads emails with an
    .xlsx or .docx attachment')."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        return read_excel_ddq(path)
    if path.suffix.lower() == ".docx":
        return read_word_ddq(path)
    raise NotImplementedError(f"No DDQ reader for {path.suffix!r} files.")
