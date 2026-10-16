"""Builds the reply to the investor: always a fresh, professional Word
document (CLAUDE.md section 3/9, as redefined by the user), regardless of
the incoming DDQ's format -- never a filled-in copy of the original file.

Built entirely from already-decided content (question text, the human's
final answer, cited sources, route) -- no model call here, matching
email_template.py's "generated from a template, not by the model". Never
writes anything from scoring/guardrails/review internals into the
document (CLAUDE.md section "Never do" extended by the user: "no routes,
scores, guardrails or reviewer names" in the investor-facing reply) --
`route` is read only to decide the Answered/Information-not-available
status, never rendered as text.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

import email_template
import scoring

BURGUNDY = RGBColor(0x8D, 0x00, 0x1C)
GREY = RGBColor(0xAD, 0xAD, 0xAD)
DARK_TEXT = RGBColor(0x2B, 0x1A, 0x1A)
FONT_NAME = "Arial"
LOGO_PATH = Path("assets/alken_fund_logo_cropped.png")

# Matches an inline citation marker in either style seen in real drafts --
# "[S11]" and "(S12)" -- including a multi-source form ("[S11, S20]"), with
# the preceding space so removing it doesn't leave a double space.
_CITATION_PATTERN = re.compile(r"\s*[\[\(]S\d+(?:\s*,\s*S\d+)*[\]\)]")


def _strip_citations(text: str) -> str:
    return _CITATION_PATTERN.sub("", text).strip()


def _style_run(run, size: int = 11, bold: bool = False, italic: bool = False,
                color: RGBColor = DARK_TEXT, name: str = FONT_NAME) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    # Word can fall back to a different font for some glyphs unless the
    # east-Asian slot is also set -- a common python-docx gotcha.
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def _add_field(paragraph, field_code: str) -> None:
    """Inserts a Word field (e.g. PAGE, NUMPAGES) -- python-docx has no
    high-level API for this, so it's built directly in OOXML."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    _style_run(run, size=9, color=GREY)


def _set_cell_borders(cell, color: str = "ADADAD", size: int = 4) -> None:
    """Thin grey borders on all 4 sides -- python-docx has no high-level
    API for table-cell borders either."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def _set_cell_text(cell, text: str, **style_kwargs) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    _style_run(run, **style_kwargs)
    _set_cell_borders(cell)


def _page_setup(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Mm(297)  # A4
    section.page_width = Mm(210)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(20)

    style = document.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(11)
    style.font.color.rgb = DARK_TEXT


def _header_footer(document: Document) -> None:
    section = document.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(header_p.add_run(f"{email_template.FUND_NAME} — DDQ Response"), size=9, color=GREY)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(footer_p.add_run("Confidential    Page "), size=9, color=GREY)
    _add_field(footer_p, "PAGE")
    _style_run(footer_p.add_run(" of "), size=9, color=GREY)
    _add_field(footer_p, "NUMPAGES")


def _cover_page(document: Document, investor_name: str, response_date: date) -> None:
    if LOGO_PATH.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Mm(55))

    for _ in range(3):
        document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(email_template.FUND_NAME), size=24, bold=True, color=BURGUNDY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(
        p.add_run(f"Response to the Due Diligence Questionnaire of {investor_name}"),
        size=15, color=BURGUNDY,
    )

    document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(f"{response_date.day} {response_date.strftime('%B %Y')}"), size=12)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(f"Information as at {email_template.MEMORANDUM_AS_OF}"), size=11)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(f"Confidential — prepared for {investor_name}"), size=11, bold=True)

    for _ in range(2):
        document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run("Fictional case study — simulated data"), size=9, color=GREY)

    document.add_page_break()


def _intro(document: Document) -> None:
    p = document.add_paragraph()
    _style_run(
        p.add_run(
            f"The answers in this document are based on the {email_template.FUND_NAME}'s "
            f"investment memorandum (version {email_template.MEMORANDUM_VERSION}, information as "
            f"at {email_template.MEMORANDUM_AS_OF}). This document is not an offer or solicitation "
            f"to invest in the fund. Where information is not available, this is stated explicitly "
            f"below."
        ),
        size=10,
    )
    document.add_paragraph()


def _status_for(question) -> str:
    """'Answered' or 'Information not available' -- never the internal
    route (AUTO/DRAFT/GAP/HUMAN_ONLY) itself. GAP means the memorandum
    doesn't cover it (by definition); a blank final answer -- however it
    got that way -- must never be presented as 'Answered' either."""
    if question.route == scoring.GAP or not (question.final_answer or "").strip():
        return "Information not available"
    return "Answered"


def _summary_table(document: Document, questions: list) -> None:
    p = document.add_paragraph()
    _style_run(p.add_run("Summary"), size=14, bold=True, color=BURGUNDY)

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    widths = (Mm(20), Mm(115), Mm(40))
    headers = table.rows[0].cells
    for cell, label, width in zip(headers, ["Ref", "Question", "Status"], widths):
        _set_cell_text(cell, label, size=10, bold=True, color=BURGUNDY)
        cell.width = width

    for question in questions:
        row = table.add_row().cells
        for cell, text, width in zip(
            row, [question.id, question.text, _status_for(question)], widths,
        ):
            _set_cell_text(cell, text, size=10)
            cell.width = width

    document.add_page_break()


def _group_by_section(questions: list) -> list[tuple[str, list]]:
    """Preserves the DDQ's own order -- both of the sections (by first
    appearance) and of the questions within each."""
    sections: dict[str, list] = {}
    order: list[str] = []
    for question in questions:
        key = question.section or "General"
        if key not in sections:
            sections[key] = []
            order.append(key)
        sections[key].append(question)
    return [(key, sections[key]) for key in order]


def _question_block(document: Document, question) -> None:
    p = document.add_paragraph()
    _style_run(p.add_run(f"{question.id} — {question.text}"), size=11, bold=True)

    answer_text = _strip_citations(question.final_answer or "")
    p = document.add_paragraph()
    _style_run(p.add_run(answer_text or "Information not available."), size=11)

    if question.cited_sources:
        p = document.add_paragraph()
        sources = ", ".join(question.cited_sources)
        _style_run(p.add_run(f"Source: Investment Memorandum, {sources}"), size=9, italic=True, color=GREY)

    document.add_paragraph()


def write_word_response(
    questions: list,
    investor_name: str,
    output_path: str | Path,
    response_date: Optional[date] = None,
) -> Path:
    """Builds the full reply document and saves it to output_path.
    `questions` is any list of objects with .id, .section, .text,
    .final_answer, .cited_sources and .route (store.QuestionState, in
    practice) -- duck-typed rather than importing store, to keep this
    module's only real dependency on the project's email/scoring
    constants."""
    response_date = response_date or date.today()
    output_path = Path(output_path)

    document = Document()
    _page_setup(document)
    _header_footer(document)
    _cover_page(document, investor_name, response_date)
    _intro(document)
    _summary_table(document, questions)

    for section_name, section_questions in _group_by_section(questions):
        p = document.add_paragraph()
        _style_run(p.add_run(section_name), size=13, bold=True, color=BURGUNDY)
        for question in section_questions:
            _question_block(document, question)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
