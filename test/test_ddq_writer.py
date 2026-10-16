"""Tests for ddq_writer.py: the reply is always a fresh Word document,
built from already-decided content only -- these tests are the guarantee
that no internal detail (route, score, guardrail name, reviewer name)
ever leaks into the investor-facing file, and that citation markers are
stripped from the visible answer text.
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx

import ddq_writer
import scoring
import store

OUT_DIR = Path(__file__).resolve().parent / "_tmp_docx_output"


def _cleanup():
    import shutil
    shutil.rmtree(OUT_DIR, ignore_errors=True)


def _question(**overrides) -> store.QuestionState:
    base = dict(
        id="Q01", section="A. Operations", text="Who is the depositary of the fund?", row=1,
        route=scoring.AUTO, final_answer="The depositary is Fictional Custody Bank SA [S13].",
        cited_sources=["S13"], approved=True,
    )
    base.update(overrides)
    return store.QuestionState(**base)


def _all_text(document) -> str:
    return "\n".join(p.text for p in document.paragraphs)


# ── citation stripping ──────────────────────────────────────────────────


def test_strip_citations_removes_square_brackets():
    assert ddq_writer._strip_citations("The fee is 0.85% [S11].") == "The fee is 0.85%."


def test_strip_citations_removes_parentheses():
    assert ddq_writer._strip_citations("Paid within five business days (S12).") == \
        "Paid within five business days."


def test_strip_citations_removes_multiple_markers():
    text = "The return was 7.65% [S10]. Monthly returns ranged widely [S10]."
    assert "[S10]" not in ddq_writer._strip_citations(text)
    assert ddq_writer._strip_citations(text).count("S10") == 0


def test_strip_citations_handles_a_multi_source_marker():
    assert ddq_writer._strip_citations("Some fact [S11, S20].") == "Some fact."


def test_strip_citations_leaves_no_double_space():
    stripped = ddq_writer._strip_citations("A fact [S11]. Another fact.")
    assert "  " not in stripped


# ── status mapping ───────────────────────────────────────────────────────


def test_status_gap_is_information_not_available():
    q = _question(route=scoring.GAP, final_answer="Partly covered.")
    assert ddq_writer._status_for(q) == "Information not available"


def test_status_empty_answer_is_information_not_available_even_if_not_gap():
    q = _question(route=scoring.HUMAN_ONLY, final_answer="")
    assert ddq_writer._status_for(q) == "Information not available"


def test_status_auto_draft_human_only_with_content_is_answered():
    for route in (scoring.AUTO, scoring.DRAFT, scoring.HUMAN_ONLY):
        q = _question(route=route, final_answer="A real answer.")
        assert ddq_writer._status_for(q) == "Answered"


# ── section grouping ─────────────────────────────────────────────────────


def test_group_by_section_preserves_first_appearance_order():
    questions = [
        _question(id="Q01", section="B. Fees"),
        _question(id="Q02", section="A. Operations"),
        _question(id="Q03", section="B. Fees"),
    ]
    grouped = ddq_writer._group_by_section(questions)
    assert [name for name, _ in grouped] == ["B. Fees", "A. Operations"]
    assert [q.id for q in grouped[0][1]] == ["Q01", "Q03"]


def test_group_by_section_handles_missing_section():
    questions = [_question(id="Q01", section=None)]
    grouped = ddq_writer._group_by_section(questions)
    assert grouped[0][0] == "General"


# ── full document build: content and no leaked internals ──────────────────


def test_full_document_has_the_expected_cover_page_text():
    _cleanup()
    try:
        questions = [_question()]
        path = ddq_writer.write_word_response(
            questions, "Harbourline Family Office", OUT_DIR / "r.docx", response_date=date(2026, 10, 16),
        )
        document = docx.Document(str(path))
        text = _all_text(document)
        assert "Alken Aurora European Equity Fund" in text
        assert "Response to the Due Diligence Questionnaire of Harbourline Family Office" in text
        assert "16 October 2026" in text
        assert "Information as at 30 September 2026" in text
        assert "Confidential — prepared for Harbourline Family Office" in text
        assert "Fictional case study — simulated data" in text
    finally:
        _cleanup()


def test_full_document_never_leaks_internal_information():
    _cleanup()
    try:
        questions = [
            _question(id="Q01", route=scoring.AUTO, final_answer="Answer one [S13]."),
            _question(id="Q02", route=scoring.DRAFT, final_answer="Answer two (S11)."),
            _question(id="Q03", route=scoring.GAP, final_answer=""),
            _question(id="Q04", route=scoring.HUMAN_ONLY, final_answer="Answer four [S15]."),
        ]
        path = ddq_writer.write_word_response(
            questions, "Harbourline Family Office", OUT_DIR / "r.docx", response_date=date(2026, 10, 16),
        )
        document = docx.Document(str(path))
        text = _all_text(document)
        for forbidden in ("AUTO", "DRAFT", "GAP", "HUMAN_ONLY", "Demo Reviewer", "guardrail", "Guardrail"):
            assert forbidden not in text, f"leaked internal term: {forbidden!r}"
        # also check table cells, not just paragraphs
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for forbidden in ("AUTO", "DRAFT", "GAP", "HUMAN_ONLY"):
                        assert forbidden not in cell.text, f"leaked internal term in table: {forbidden!r}"
    finally:
        _cleanup()


def test_full_document_has_summary_table_with_ref_question_status():
    _cleanup()
    try:
        questions = [
            _question(id="Q01", route=scoring.AUTO, final_answer="Answered one."),
            _question(id="Q02", route=scoring.GAP, final_answer=""),
        ]
        path = ddq_writer.write_word_response(
            questions, "Harbourline Family Office", OUT_DIR / "r.docx", response_date=date(2026, 10, 16),
        )
        document = docx.Document(str(path))
        table = document.tables[0]
        assert [c.text for c in table.rows[0].cells] == ["Ref", "Question", "Status"]
        assert table.rows[1].cells[0].text == "Q01"
        assert table.rows[1].cells[2].text == "Answered"
        assert table.rows[2].cells[2].text == "Information not available"
    finally:
        _cleanup()


def test_full_document_strips_citations_from_visible_answer_text():
    _cleanup()
    try:
        questions = [_question(final_answer="The fee is 0.85% [S11].", cited_sources=["S11"])]
        path = ddq_writer.write_word_response(
            questions, "Harbourline Family Office", OUT_DIR / "r.docx", response_date=date(2026, 10, 16),
        )
        document = docx.Document(str(path))
        text = _all_text(document)
        assert "[S11]" not in text
        assert "The fee is 0.85%." in text
        assert "Source: Investment Memorandum, S11" in text
    finally:
        _cleanup()


def test_full_document_header_and_footer():
    _cleanup()
    try:
        path = ddq_writer.write_word_response(
            [_question()], "Harbourline Family Office", OUT_DIR / "r.docx", response_date=date(2026, 10, 16),
        )
        document = docx.Document(str(path))
        assert document.sections[0].header.paragraphs[0].text == "Alken Aurora European Equity Fund — DDQ Response"
        assert "Confidential" in document.sections[0].footer.paragraphs[0].text
    finally:
        _cleanup()


def test_full_document_is_a4_with_generous_margins():
    _cleanup()
    try:
        path = ddq_writer.write_word_response(
            [_question()], "Harbourline Family Office", OUT_DIR / "r.docx", response_date=date(2026, 10, 16),
        )
        document = docx.Document(str(path))
        section = document.sections[0]
        assert round(section.page_width.mm) == 210
        assert round(section.page_height.mm) == 297
        assert section.left_margin.mm >= 20
    finally:
        _cleanup()
